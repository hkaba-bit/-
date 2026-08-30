#!/usr/bin/env python3
"""
生成した仕様書サイトマップシートを、品質チェックリストで検証する。

  python verify_sitemap_xlsx.py 仕様書_クライアント名_提案サイトマップ.xlsx

recalc.py を通した後に実行すること（集計値の検証に必要）。
WARN は必ずしも誤りではない。案件の事情で意図的にそうしている場合もあるので、
指摘された行を見て、意図どおりならそのまま進めてよい。
"""

import sys
import openpyxl

MARK = "■"
UPD_VALID = {"◎", "○", "△", "×"}


def find_data_range(ws):
    """ヘッダー3段の下からデータ行を探し、集計行の手前までを返す。"""
    start = 7
    last = start - 1
    for r in range(start, ws.max_row + 1):
        b = ws.cell(row=r, column=2).value
        if b and "合計" in str(b):
            break
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, 7)):
            last = r
    return start, last


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    start, last = find_data_range(ws)
    if last < start:
        print("FAIL  データ行が見つかりません")
        return 1

    # 工程列の範囲を特定（更新性列の手前まで）
    col_upd = None
    for c in range(7, ws.max_column + 1):
        if ws.cell(row=4, column=c).value == "更新性":
            col_upd = c
            break
    if col_upd is None:
        print("FAIL  更新性の列が見つかりません")
        return 1
    proc0, col_memo = 7, col_upd + 1

    issues, warns = [], []

    # 第1階層の項目数
    l1 = [ws.cell(row=r, column=2).value for r in range(start, last + 1)]
    l1 = [v for v in l1 if v]
    if len(l1) > 6:
        warns.append(
            f"第1階層が{len(l1)}項目: {', '.join(map(str, l1))}\n"
            "      → TOP・詳細ページ・CV導線・規約類を除いた"
            "「グローバルナビに載る項目」で数え直し、6以内か確認する")

    # URL・更新性・工程マーク
    no_url, bad_upd, all_marked, no_memo = [], [], [], []
    for r in range(start, last + 1):
        if not ws.cell(row=r, column=6).value:
            no_url.append(r)
        upd = ws.cell(row=r, column=col_upd).value
        if upd not in UPD_VALID:
            bad_upd.append(r)
        marks = [ws.cell(row=r, column=c).value for c in range(proc0, col_upd)]
        if all(m == MARK for m in marks):
            all_marked.append(r)
        if not ws.cell(row=r, column=col_memo).value:
            no_memo.append(r)

    if no_url:
        issues.append(f"URLが空の行: {no_url}")
    if bad_upd:
        issues.append(f"更新性が未記入または不正な行: {bad_upd}")
    if all_marked:
        issues.append(f"全工程に■が付いている行（見積が読めない）: {all_marked}")

    n = last - start + 1
    if len(no_memo) > n / 2:
        warns.append(f"備考が空の行が{len(no_memo)}/{n}行（半数以下を推奨）")

    # 更新性◎○に更新手段があるか
    hdr = {c: ws.cell(row=6, column=c).value for c in range(proc0, col_upd)}
    wp_col = next((c for c, v in hdr.items() if v == "WP"), None)
    if wp_col:
        for r in range(start, last + 1):
            if ws.cell(row=r, column=col_upd).value in ("◎", "○"):
                if ws.cell(row=r, column=wp_col).value != MARK:
                    memo = ws.cell(row=r, column=col_memo).value or ""
                    warns.append(
                        f"{r}行目は更新性{ws.cell(row=r, column=col_upd).value}だが"
                        f"WP列が空。更新手段を確認（備考: {str(memo)[:30]}）")

    # 集計行
    sum_row = None
    for r in range(last + 1, min(last + 5, ws.max_row) + 1):
        b = ws.cell(row=r, column=2).value
        if b and "合計" in str(b):
            sum_row = r
            break
    if sum_row is None:
        issues.append("集計行（合計）が見つかりません")
    else:
        vals = [ws.cell(row=sum_row, column=c).value for c in range(proc0, col_upd)]
        if any(v is None for v in vals):
            issues.append("集計行にキャッシュ値がありません。recalc.py を実行してください")
        elif all(v == 0 for v in vals):
            issues.append("集計値がすべて0です。■が入っているか確認してください")

    # 付帯ブロック
    text = " ".join(
        str(ws.cell(row=r, column=c).value or "")
        for r in range(last + 1, ws.max_row + 1) for c in (2, 4))
    for key, label in [("凡例", "凡例"), ("対象外", "制作対象外ブロック"),
                       ("301", "URL移行ブロック（301の記載）"),
                       ("要確認", "要確認事項ブロック")]:
        if key not in text:
            warns.append(f"{label}が見つかりません")

    print(f"シート: {ws.title} ／ データ {n}行（{start}〜{last}）")
    print()
    for m in issues:
        print(f"FAIL  {m}")
    for m in warns:
        print(f"WARN  {m}")
    if not issues and not warns:
        print("OK    チェック項目をすべて満たしています")
    elif not issues:
        print()
        print("FAILなし。WARNは内容を確認のうえ、意図どおりなら問題ありません。")
    return 1 if issues else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("使い方: python verify_sitemap_xlsx.py <xlsx>")
    sys.exit(main(sys.argv[1]))
