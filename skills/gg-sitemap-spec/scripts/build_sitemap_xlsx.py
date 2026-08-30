#!/usr/bin/env python3
"""
仕様書サイトマップシートのビルダー。

JSON定義を読み、GrowGroupの仕様書フォーマットに沿ったxlsxを生成する。

  python build_sitemap_xlsx.py spec.json -o 仕様書_クライアント名_提案サイトマップ.xlsx

生成後は必ず recalc.py を通すこと（openpyxlは数式のキャッシュ値を持たない）:

  python /mnt/skills/public/xlsx/scripts/recalc.py <出力ファイル>

JSONの書き方は assets/sitemap_spec_example.json を参照。
"""

import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "ＭＳ ゴシック"
MARK = "■"

DEFAULT_COLUMNS = [
    {"group": "構成", "subs": ["TOP", "SP", "下層A", "下層B"]},
    {"group": "デザイン", "subs": ["TOP", "SP", "下層", "フォーマット"]},
    {"group": "コーディング", "subs": ["TOP", "SP", "下層A", "下層B", "WP", "フォーム"]},
]

DEFAULT_LEGEND = [
    ("■", "制作対象。該当する工程にマークを入れる"),
    ("更新性 ◎", "公開後に頻繁に更新（WordPressで運用）"),
    ("更新性 ○", "年数回の更新（WordPressで運用）"),
    ("更新性 △", "ほぼ更新しない（固定ページ）"),
    ("更新性 ×", "更新しない"),
    ("下層A／下層B", "下層ページの制作区分。単価区分は自社の単価表に合わせて置き換える"),
]

UPD_COLOR = {"◎": "C00000", "○": "0070C0"}

thin = Side(style="thin", color="808080")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
GRP_FILL = PatternFill("solid", fgColor="BFBFBF")
L1_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")


def build(spec: dict, out_path: str) -> None:
    columns = spec.get("columns") or DEFAULT_COLUMNS
    rows = spec.get("rows") or []
    if not rows:
        sys.exit("エラー: rows が空です。1行以上定義してください。")

    n_proc = sum(len(g["subs"]) for g in columns)
    col_url = 6
    col_proc0 = 7
    col_upd = col_proc0 + n_proc
    col_memo = col_upd + 1
    last_col = col_memo

    wb = Workbook()
    ws = wb.active
    ws.title = spec.get("sheet_name", "提案サイトマップ")[:31]

    def cell(r, c, v=None, bold=False, size=9, fill=None,
             align="left", wrap=False, color=None):
        x = ws.cell(row=r, column=c)
        if v is not None:
            x.value = v
        x.font = Font(name=FONT, size=size, bold=bold, color=color)
        x.border = BOX
        x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fill:
            x.fill = fill
        return x

    def plain(r, c, v, bold=False, size=8):
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name=FONT, size=size, bold=bold)
        return x

    # ---- タイトル ----
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    t = ws.cell(row=1, column=1, value=spec.get("title", "提案 サイトマップ"))
    t.font = Font(name=FONT, size=12, bold=True)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    s = ws.cell(row=2, column=1, value=spec.get("subtitle", ""))
    s.font = Font(name=FONT, size=8)
    s.alignment = Alignment(horizontal="right", vertical="center")

    # ---- ヘッダー3段 ----
    HR1, HR2, HR3 = 4, 5, 6
    ws.merge_cells(start_row=HR1, start_column=1, end_row=HR1, end_column=col_url)
    cell(HR1, 1, "論理構造", bold=True, fill=GRP_FILL, align="center")
    ws.merge_cells(start_row=HR1, start_column=col_proc0,
                   end_row=HR1, end_column=col_proc0 + n_proc - 1)
    cell(HR1, col_proc0, "見積", bold=True, fill=GRP_FILL, align="center")
    ws.merge_cells(start_row=HR1, start_column=col_upd, end_row=HR3, end_column=col_upd)
    cell(HR1, col_upd, "更新性", bold=True, fill=GRP_FILL, align="center")
    ws.merge_cells(start_row=HR1, start_column=col_memo, end_row=HR3, end_column=col_memo)
    cell(HR1, col_memo, "備考・根拠", bold=True, fill=GRP_FILL, align="center")

    for i, name in enumerate(["ID", "第1階層", "第2階層", "第3階層", "第4階層", "URL"]):
        ws.merge_cells(start_row=HR2, start_column=1 + i, end_row=HR3, end_column=1 + i)
        cell(HR2, 1 + i, name, bold=True, fill=HDR_FILL, align="center")

    c = col_proc0
    for g in columns:
        subs = g["subs"]
        ws.merge_cells(start_row=HR2, start_column=c, end_row=HR2, end_column=c + len(subs) - 1)
        cell(HR2, c, g["group"], bold=True, fill=HDR_FILL, align="center")
        for j, sname in enumerate(subs):
            cell(HR3, c + j, sname, bold=True, fill=HDR_FILL, align="center", size=8)
        c += len(subs)

    for r in (HR1, HR2, HR3):
        ws.row_dimensions[r].height = 18

    # ---- データ行 ----
    START = 7
    r = START
    for row in rows:
        l1 = row.get("l1", "")
        is_l1 = bool(l1)
        fill = L1_FILL if is_l1 else None

        cell(r, 1, row.get("id", ""), align="center", fill=fill)
        cell(r, 2, l1, bold=is_l1, fill=fill)
        cell(r, 3, row.get("l2", ""), fill=fill)
        cell(r, 4, row.get("l3", ""), fill=fill)
        cell(r, 5, row.get("l4", ""), fill=fill)
        cell(r, col_url, row.get("url", ""), fill=fill, size=8)

        marks = row.get("marks", {})
        c = col_proc0
        for g in columns:
            got = marks.get(g["group"], [])
            for sname in g["subs"]:
                v = MARK if sname in got else ""
                cell(r, c, v, align="center", fill=fill)
                c += 1

        upd = row.get("upd", "")
        cell(r, col_upd, upd, align="center", fill=fill, bold=True,
             color=UPD_COLOR.get(upd))

        memo = row.get("memo", "")
        cell(r, col_memo, memo, size=8, wrap=True)
        ws.row_dimensions[r].height = 30 if memo else 18
        r += 1

    LAST = r - 1

    # ---- 集計行 ----
    r += 1
    cell(r, 2, "合計（■の数）", bold=True, fill=HDR_FILL)
    for c in range(col_proc0, col_upd):
        col = get_column_letter(c)
        cell(r, c, f'=COUNTIF({col}{START}:{col}{LAST},"■")',
             align="center", bold=True, fill=HDR_FILL)

    # ---- 凡例 ----
    r += 2
    plain(r, 2, "凡例", bold=True, size=10)
    r += 1
    for k, v in spec.get("legend", DEFAULT_LEGEND):
        plain(r, 2, k, bold=True)
        plain(r, 4, v)
        r += 1

    # ---- 付帯ブロック ----
    for block_key, heading in [
        ("excluded", "今回は制作対象外（現状のまま据え置き）"),
        ("migration", "URL移行"),
    ]:
        items = spec.get(block_key)
        if not items:
            continue
        r += 1
        h = cell(r, 2, heading, bold=True, size=10, fill=NOTE_FILL)
        h.border = Border()
        r += 1
        for k, v in items:
            plain(r, 2, k, bold=True)
            plain(r, 4, v)
            r += 1

    todo = spec.get("todo")
    if todo:
        r += 1
        h = cell(r, 2, "要確認事項", bold=True, size=10, fill=NOTE_FILL)
        h.border = Border()
        r += 1
        for i, v in enumerate(todo, 1):
            plain(r, 2, f"{i}.", bold=True)
            plain(r, 4, v)
            r += 1

    # ---- 体裁 ----
    widths = {1: 5, 2: 16, 3: 18, 4: 16, 5: 12, col_url: 24,
              col_upd: 7, col_memo: 60}
    for c in range(col_proc0, col_upd):
        widths[c] = 6
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = f"{get_column_letter(col_proc0)}{START}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f"保存: {out_path}")
    print(f"データ行 {START}〜{LAST}（{LAST - START + 1}行）／工程列 {n_proc}")
    print("次に recalc.py を実行してください。")


def main():
    ap = argparse.ArgumentParser(description="仕様書サイトマップシートを生成する")
    ap.add_argument("spec", help="JSON定義ファイル")
    ap.add_argument("-o", "--out", required=True, help="出力先の .xlsx")
    a = ap.parse_args()
    with open(a.spec, encoding="utf-8") as f:
        build(json.load(f), a.out)


if __name__ == "__main__":
    main()
