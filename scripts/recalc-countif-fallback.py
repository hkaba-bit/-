#!/usr/bin/env python3
"""
COUNTIF のキャッシュ値を openpyxl だけで埋める代替スクリプト。

`xlsx` Skill の `recalc.py` が使えない環境向けの**フォールバック**。
LibreOffice が動く環境では必ず `recalc.py` を使うこと（こちらは
`=COUNTIF(範囲,"文字列")` しか評価できない）。

  python scripts/recalc-countif-fallback.py <xlsx>

背景：Claude Code on the web の実行環境では soffice が xlsx を読み込めず
（"source file could not be loaded"）、recalc.py がタイムアウトする。
仕様書サイトマップの集計行は COUNTIF のみなので、値を計算して
sheet XML の <c> に <v> を差し込む。
"""

import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

COUNTIF_RE = re.compile(r'^=COUNTIF\(([A-Z]+)(\d+):([A-Z]+)(\d+),"([^"]*)"\)$')


def main(path: str) -> int:
    src = Path(path)
    wb = openpyxl.load_workbook(src)  # 数式そのものを読む

    # シートごとに {セル座標: 計算値}
    computed: dict[str, dict[str, int]] = {}
    total = 0
    for ws in wb.worksheets:
        hits: dict[str, int] = {}
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                m = COUNTIF_RE.match(c.value)
                if not m:
                    continue
                c0, r0, c1, r1, crit = m.groups()
                n = 0
                for rr in range(int(r0), int(r1) + 1):
                    for cc in range(column_index_from_string(c0),
                                    column_index_from_string(c1) + 1):
                        if ws.cell(row=rr, column=cc).value == crit:
                            n += 1
                hits[c.coordinate] = n
                total += 1
        if hits:
            computed[ws.title] = hits

    if not total:
        print("COUNTIF の数式が見つかりませんでした")
        return 1

    # sheet XML に <v> を差し込む
    sheet_paths = {}
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        for i, ws in enumerate(wb.worksheets, start=1):
            p = f"xl/worksheets/sheet{i}.xml"
            if ws.title in computed and p in names:
                sheet_paths[p] = computed[ws.title]
        payload = {n: z.read(n) for n in names}

    patched = 0
    for p, hits in sheet_paths.items():
        xml = payload[p].decode("utf-8")
        for coord, val in hits.items():
            pat = re.compile(
                r'(<c r="%s"(?:\s[^>]*?)?)(/>|>)(<f>[^<]*</f>)?(?:<v>[^<]*</v>)?'
                r'(?:</c>)?' % re.escape(coord))

            def repl(m, val=val):
                attrs, close, f = m.group(1), m.group(2), m.group(3) or ""
                attrs = re.sub(r'\st="[^"]*"', "", attrs)
                if close == "/>":
                    return f'{attrs} t="n"/>'
                return f'{attrs} t="n">{f}<v>{val}</v></c>'

            xml, n = pat.subn(repl, xml, count=1)
            patched += n
        payload[p] = xml.encode("utf-8")

    tmp = src.with_suffix(".xlsx.tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in payload.items():
            z.writestr(n, data)
    shutil.move(str(tmp), str(src))

    print(f"COUNTIF {total}件を計算し、{patched}件にキャッシュ値を書き込みました: {src}")
    return 0 if patched == total else 1


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("使い方: python scripts/recalc-countif-fallback.py <xlsx>")
    sys.exit(main(sys.argv[1]))
